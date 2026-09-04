"""测试PPT生成流程"""
import sys
import os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from app.services.ppt_content_generator import PPTContentGenerator

print("=== PPT生成测试 ===")

# 1. 测试内容生成
print("\n1. 生成PPT内容...")
g = PPTContentGenerator()
content = g.generate_ppt_content(
    topic='细胞分裂',
    duration=45,
    audience='高中生',
    subject='生物学',
    theme_color='蓝色系',
    style_preference='现代简洁',
    key_points=['有丝分裂', '减数分裂']
)
print(f"   ✓ 大纲页数: {len(content.get('outline', []))}")
print(f"   ✓ 详细内容页数: {len(content.get('detailed_content', []))}")

# 打印第一页内容看看格式
if content.get('detailed_content'):
    first_page = content['detailed_content'][0]
    print(f"   第一页类型: {first_page.get('page_type')}")
    print(f"   第一页内容键: {list(first_page.keys())}")

# 保存内容到JSON，方便调试
import json
with open('debug_ppt_content.json', 'w', encoding='utf-8') as f:
    json.dump(content, f, ensure_ascii=False, indent=2)
print("   ✓ 内容已保存到 debug_ppt_content.json")

# 2. 测试PPT文件生成 - 分步测试
print("\n2. 分步测试PPT生成...")

# 先测试基本的python-pptx功能
print("   测试基本Presentation创建...")
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    prs.save('test_basic.pptx')
    print("   ✓ 基本Presentation创建成功")
except Exception as e:
    print(f"   ✗ 基本Presentation创建失败: {e}")
    import traceback
    traceback.print_exc()

# 现在测试完整的PPT生成
print("\n3. 测试完整PPT生成...")
from app.services.ppt_file_generator import PPTFileGenerator
f = PPTFileGenerator()

try:
    pptx_path = f.generate_pptx(content)
    print(f"   ✓ PPT文件生成成功: {pptx_path}")
except Exception as e:
    print(f"   ✗ 失败: {e}")
    import traceback
    traceback.print_exc()

print("\n=== 测试完成 ===")
