"""
A04 教学资源生成：动画创意、互动游戏、完整教案（LLM + 离线模板回退）。
赛题方向：多模态 AI 辅助教学 — 需可运行的 HTML5 预览与规范教案结构。
"""
from __future__ import annotations

import html
import json
import re
import uuid
from pathlib import Path
from typing import Any, Dict, List, Optional

from . import llm_client
from .paths import GENERATED_ANIM, GENERATED_GAMES, GENERATED_CW

ANIMATION_STYLES: List[Dict[str, str]] = [
    {"id": "2D动画", "name": "2D 扁平插画"},
    {"id": "3D风格", "name": "3D 轻质感"},
    {"id": "白板手绘", "name": "白板手绘风"},
    {"id": "科技风", "name": "科技 / 数据可视化"},
    {"id": "国风", "name": "中国风 / 水墨点缀"},
    {"id": "简约", "name": "极简信息图"},
]

GAME_TYPES: List[Dict[str, str]] = [
    {"id": "quiz", "name": "单选题闯关"},
    {"id": "flashcard", "name": "记忆翻卡"},
    {"id": "match", "name": "连线配对"},
    {"id": "sort", "name": "排序挑战"},
    {"id": "fill", "name": "填空测验"},
]


def _has_real_llm() -> bool:
    import os

    from . import persist

    s = persist.get_active_llm_settings()
    if s.get("provider") == "demo":
        return False
    if os.getenv("DASHSCOPE_API_KEY") or os.getenv("OPENAI_API_KEY"):
        return True
    return bool(s.get("api_key"))


def _normalize_game_type(raw: str) -> str:
    r = (raw or "quiz").strip()
    m = {
        "选择题": "quiz",
        "单选题": "quiz",
        "问答闯关": "quiz",
        "翻卡": "flashcard",
        "记忆翻卡": "flashcard",
        "连线": "match",
        "配对": "match",
        "排序": "sort",
        "填空": "fill",
    }
    known = {"quiz", "flashcard", "match", "sort", "fill"}
    return m.get(r, r if r in known else "quiz")


def _extract_html_block(text: str) -> str:
    """Pull first <!DOCTYPE html> or <html ...> ... </html> from model output."""
    if not text:
        return ""
    t = text.strip()
    m = re.search(r"(<!DOCTYPE\s+html[\s\S]*?</html>)", t, re.I)
    if m:
        return m.group(1).strip()
    m = re.search(r"(<html[\s\S]*?</html>)", t, re.I)
    if m:
        return m.group(1).strip()
    if "```" in t:
        inner = re.sub(r"^```\w*\n?", "", t)
        inner = re.sub(r"\n?```$", "", inner, flags=re.M)
        return inner.strip()
    return t


def _demo_animation_html(topic: str, audience: str, teaching: str, style: str, duration: int) -> str:
    esc = html.escape
    return f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8"/>
<meta name="viewport" content="width=device-width, initial-scale=1"/>
<title>{esc(topic)} — 教学动画预览</title>
<style>
  * {{ box-sizing: border-box; }}
  body {{
    margin: 0; min-height: 100vh; font-family: "Segoe UI", "Microsoft YaHei", sans-serif;
    background: linear-gradient(135deg, #0f172a 0%, #1e3a5f 50%, #0f172a 100%);
    color: #e2e8f0; overflow-x: hidden;
  }}
  .stage {{
    max-width: 920px; margin: 0 auto; padding: 24px;
  }}
  h1 {{
    font-size: 1.6rem; text-align: center;
    background: linear-gradient(90deg, #38bdf8, #a78bfa);
    -webkit-background-clip: text; -webkit-text-fill-color: transparent;
    margin-bottom: 8px;
  }}
  .meta {{ text-align: center; color: #94a3b8; font-size: 0.9rem; margin-bottom: 24px; }}
  .scene {{
    position: relative; height: 320px; border-radius: 16px;
    background: rgba(15,23,42,0.6); border: 1px solid rgba(56,189,248,0.3);
    overflow: hidden; margin-bottom: 20px;
  }}
  .orb {{
    position: absolute; width: 80px; height: 80px; border-radius: 50%;
    background: radial-gradient(circle at 30% 30%, #38bdf8, #6366f1);
    box-shadow: 0 0 40px rgba(56,189,248,0.5);
    animation: float 4s ease-in-out infinite;
  }}
  .orb.b {{ left: 15%; top: 40%; animation-delay: 0.5s; }}
  .orb.c {{ right: 18%; top: 28%; animation-delay: 1s; background: radial-gradient(circle at 30% 30%, #f472b6, #a78bfa); }}
  @keyframes float {{
    0%,100% {{ transform: translate(0,0) scale(1); }}
    50% {{ transform: translate(20px,-24px) scale(1.05); }}
  }}
  .caption {{
    position: absolute; bottom: 0; left: 0; right: 0; padding: 14px 18px;
    background: linear-gradient(transparent, rgba(15,23,42,0.95));
    font-size: 0.95rem; line-height: 1.5;
  }}
  .story {{
    display: grid; gap: 10px;
  }}
  .shot {{
    padding: 12px 14px; border-radius: 10px;
    background: rgba(30,41,59,0.7); border-left: 3px solid #38bdf8;
    font-size: 0.9rem;
  }}
</style>
</head>
<body>
  <div class="stage">
    <h1>{esc(topic)}</h1>
    <div class="meta">受众：{esc(audience)} ｜ 风格：{esc(style)} ｜ 建议时长约 {duration} 分钟</div>
    <div class="scene">
      <div class="orb" style="left:42%;top:30%;"></div>
      <div class="orb b"></div>
      <div class="orb c"></div>
      <div class="caption">演示动画场景：核心概念可视化 — {esc(teaching[:80])}{"…" if len(teaching)>80 else ""}</div>
    </div>
    <div class="story">
      <div class="shot"><strong>分镜1 · 导入</strong>：创设情境，引出主题「{esc(topic)}」。</div>
      <div class="shot"><strong>分镜2 · 展开</strong>：分步展示关键知识点与动态关系。</div>
      <div class="shot"><strong>分镜3 · 小结</strong>：归纳要点，衔接课堂练习。</div>
    </div>
    <p style="text-align:center;color:#64748b;font-size:0.8rem;margin-top:16px;">
      （演示模板）配置大模型 API 后将根据您的描述生成更精细的分镜与动画脚本。
    </p>
  </div>
</body>
</html>"""


def _demo_game_html(topic: str, audience: str, teaching: str, game_type: str) -> str:
    esc = html.escape
    gt = game_type or "quiz"
    qa = [
        (f"关于「{topic}」，下列哪项描述最准确？", ["核心定义与典型特征", "无关内容", "错误表述"], 0),
        (f"在「{teaching[:40]}…」情境下，第一步应关注什么？", ["目标与约束", "随机尝试", "跳过预习"], 0),
        (f"巩固题：本课重点与下列哪一条最相关？", ["知识迁移与应用", "完全无关", "逆向混淆项"], 0),
    ]
    opts_html = ""
    for qi, (qtext, opts, _) in enumerate(qa):
        opts_html += f'<div class="q" data-i="{qi}"><p class="qt">{esc(qtext)}</p><div class="opts">'
        for j, o in enumerate(opts):
            opts_html += f'<button type="button" class="opt" data-c="{1 if j==0 else 0}">{esc(o)}</button>'
        opts_html += "</div></div>"

    return f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8"/>
<meta name="viewport" content="width=device-width, initial-scale=1"/>
<title>{esc(topic)} — 互动游戏</title>
<style>
  body {{
    margin: 0; min-height: 100vh; font-family: "Microsoft YaHei", sans-serif;
    background: #0b1220; color: #e5e7eb;
  }}
  .wrap {{ max-width: 720px; margin: 0 auto; padding: 20px; }}
  h1 {{ font-size: 1.35rem; color: #38bdf8; }}
  .sub {{ color: #94a3b8; font-size: 0.85rem; margin-bottom: 16px; }}
  .q {{ display: none; background: #111827; border-radius: 12px; padding: 16px; margin-bottom: 12px; border: 1px solid #1f2937; }}
  .q.on {{ display: block; }}
  .qt {{ margin: 0 0 12px; line-height: 1.5; }}
  .opts {{ display: flex; flex-direction: column; gap: 8px; }}
  .opt {{
    text-align: left; padding: 10px 12px; border-radius: 8px; border: 1px solid #334155;
    background: #1e293b; color: #e5e7eb; cursor: pointer; font-size: 0.95rem;
  }}
  .opt:hover {{ border-color: #38bdf8; }}
  .bar {{ height: 6px; background: #1f2937; border-radius: 4px; overflow: hidden; margin: 12px 0; }}
  .fill {{ height: 100%; background: linear-gradient(90deg,#22c55e,#38bdf8); width: 0%; transition: width .3s; }}
  #msg {{ margin-top: 12px; font-weight: bold; }}
</style>
</head>
<body>
<div class="wrap">
  <h1>{esc(topic)}</h1>
  <div class="sub">类型：{esc(gt)} ｜ 受众：{esc(audience)}</div>
  <div class="bar"><div class="fill" id="prog"></div></div>
  {opts_html}
  <p id="msg"></p>
</div>
<script>
(function(){{
  var qs = [].slice.call(document.querySelectorAll('.q'));
  var i = 0, score = 0;
  function show() {{
    qs.forEach(function(q, k) {{ q.classList.toggle('on', k === i); }});
    document.getElementById('prog').style.width = (i / qs.length * 100) + '%';
  }}
  document.querySelectorAll('.opt').forEach(function(btn) {{
    btn.addEventListener('click', function() {{
      if (parseInt(btn.getAttribute('data-c'),10) === 1) score++;
      i++;
      if (i >= qs.length) {{
        document.getElementById('msg').textContent = '完成！得分：' + score + ' / ' + qs.length;
        document.getElementById('prog').style.width = '100%';
        return;
      }}
      show();
    }});
  }});
  show();
}})();
</script>
</body>
</html>"""


def generate_animation_payload(body: Dict[str, Any]) -> Dict[str, Any]:
    topic = (body.get("topic") or "课程主题").strip()
    audience = (body.get("audience") or "大学生").strip()
    teaching = (body.get("teaching_content") or topic).strip()
    style = (body.get("style") or "2D动画").strip()
    duration = int(body.get("duration") or 5)

    html_out = ""
    story_rows = [
        {"shot": 1, "title": "情境导入", "content": f"引出主题：{topic}", "duration_sec": 30},
        {"shot": 2, "title": "知识展开", "content": "核心概念动态演示与旁白要点", "duration_sec": 120},
        {"shot": 3, "title": "归纳小结", "content": "关键结论与课堂衔接", "duration_sec": 40},
    ]
    scenes_ui = [
        {
            "scene_id": r["shot"],
            "title": r["title"],
            "description": r["content"],
            "duration": f"{r['duration_sec']}秒",
        }
        for r in story_rows
    ]
    animation_data: Dict[str, Any] = {
        "title": f"{topic} — 教学动画方案",
        "description": f"面向{audience}，采用「{style}」视觉风格，围绕「{teaching[:200]}」展开分镜与动效设计。",
        "duration_minutes": duration,
        "style": style,
        "estimated_time": f"约 {duration} 分钟",
        "storyboard": story_rows,
        "scenes": scenes_ui,
        "characters": [
            {"name": "教学旁白", "role": "讲解与节奏控制"},
            {"name": "图示元素", "role": "概念可视化"},
        ],
    }

    if _has_real_llm():
        prompt = f"""你是教学动画导演与前端工程师。根据以下需求生成一份可直接在浏览器播放的单文件 HTML5 动画演示（仅 HTML+CSS+少量 JS，可内联）。
主题：{topic}
受众：{audience}
教学内容：{teaching}
动画风格：{style}
建议总时长约 {duration} 分钟体量的演示节奏（可用 CSS 动画表现节奏感）。

要求：
1. 使用中文；视觉现代、对比清晰、适合课堂投影。
2. 必须包含完整 <!DOCTYPE html> … </html>，不要 Markdown。
3. 页面中用文字区块简要列出 3~5 个分镜说明（可用 div 列表）。
4. 动画以 CSS @keyframes / transform 为主，避免外链资源。

直接输出 HTML 源码。"""
        raw = llm_client.blocking_chat([{"role": "user", "content": prompt}], max_tokens=8192)
        html_out = _extract_html_block(raw)
        if len(html_out) < 80:
            html_out = _demo_animation_html(topic, audience, teaching, style, duration)
    else:
        html_out = _demo_animation_html(topic, audience, teaching, style, duration)

    tid = "anim_" + uuid.uuid4().hex[:10]
    d = GENERATED_ANIM / tid
    d.mkdir(parents=True, exist_ok=True)
    (d / "preview.html").write_text(html_out, encoding="utf-8")

    return {
        "html_content": html_out,
        "task_dir": str(d.resolve()),
        "topic": topic,
        "animation_data": animation_data,
    }


def generate_game_payload(body: Dict[str, Any]) -> Dict[str, Any]:
    topic = (body.get("topic") or "课堂主题").strip()
    audience = (body.get("audience") or "大学生").strip()
    teaching = (body.get("teaching_content") or topic).strip()
    game_type = _normalize_game_type(body.get("game_type") or "quiz")

    questions = [
        {"question": f"下列哪项最符合「{topic}」的核心要点？", "options": ["准确定义与典型场景", "干扰项B", "干扰项C"], "answer": 0},
        {"question": f"结合「{teaching[:60]}…」，学习时应优先关注？", "options": ["目标—策略—迁移", "随机试错", "忽略反馈"], "answer": 0},
    ]

    game_data = {
        "title": f"{topic} · 互动游戏",
        "description": f"基于教学主题生成的{game_type}类课堂小游戏，受众：{audience}。",
        "game_type": game_type,
        "difficulty": "适中",
        "estimated_time": "5-10分钟",
        "rules": [
            "阅读题干后选择你认为最恰当的选项。",
            "可重复闯关，建议记录错题以便讲评。",
        ],
        "questions": [{"question": q["question"], "options": q["options"], "correct": q["answer"]} for q in questions],
    }

    if _has_real_llm():
        prompt = f"""你是教学游戏设计师与前端工程师。请为课堂生成一个单文件 HTML5 小游戏（仅内联 HTML/CSS/JS）。
主题：{topic}
受众：{audience}
教学内容要点：{teaching}
游戏类型：{game_type}（若为 quiz 则做单选题闯关；match 可做简易翻牌配对；flashcard 做翻卡；sort 做拖拽排序；fill 做填空）

要求：
1. 中文界面，风格清爽，适合课堂。
2. 完整 <!DOCTYPE html> … </html>，无 Markdown。
3. 包含简单计分或完成提示。
4. 不使用外部 CDN 图片（可用 Unicode 或 CSS 图形）。

直接输出 HTML 源码。"""
        raw = llm_client.blocking_chat([{"role": "user", "content": prompt}], max_tokens=8192)
        html_out = _extract_html_block(raw)
        if len(html_out) < 80:
            html_out = _demo_game_html(topic, audience, teaching, game_type)
    else:
        html_out = _demo_game_html(topic, audience, teaching, game_type)

    tid = "game_" + uuid.uuid4().hex[:10]
    d = GENERATED_GAMES / tid
    d.mkdir(parents=True, exist_ok=True)
    (d / "game.html").write_text(html_out, encoding="utf-8")

    return {
        "html_content": html_out,
        "task_dir": str(d.resolve()),
        "topic": topic,
        "game_data": game_data,
    }


def generate_lesson_plan_full(body: Dict[str, Any]) -> Dict[str, Any]:
    """符合常见中小学/高校教案栏目：课题、课型、目标、重难点、过程、作业、板书。"""
    topic = (body.get("topic") or body.get("subject") or "课题").strip()
    level = (body.get("level") or "普通高等院校").strip()
    grade = (body.get("grade") or "大一").strip()
    duration = int(body.get("duration") or 45)
    requirements = (body.get("requirements") or "").strip()

    skeleton = {
        "title": f"《{topic}》教学设计",
        "meta": {
            "课题": topic,
            "课程学段": level,
            "年级/班级": grade,
            "课型": "新授课",
            "课时": f"{duration} 分钟",
        },
        # 注意：这里不用"知识与技能/过程与方法/情感态度与价值观"三维目标。
        # 2022 年版课程方案与课程标准已明确"从三维目标转向课程核心素养"，
        # 兜底骨架也必须按核心素养的通用维度给，否则兜底内容本身就与现行课标冲突。
        "objectives": {
            "核心素养·观念理解": [f"理解{topic}的核心概念与基本应用。"],
            "核心素养·思维方法": ["通过情境探究与合作讨论，经历“问题—证据—结论”的学习过程。"],
            "核心素养·态度责任": ["激发探究兴趣，形成严谨求实的科学态度与合作意识。"],
            "课标依据": ["（需由教师按本学科现行课程标准补充）"],
        },
        "key_points": [f"{topic}的关键概念与典型应用"],
        "difficulties": ["易混淆点辨析与迁移应用"],
        "teaching_process": [
            {
                "stage": "一、导入新课",
                "time_min": 5,
                "teacher_activity": "创设情境，提出驱动问题，明确学习目标。",
                "student_activity": "观察思考，提出初步猜想。",
            },
            {
                "stage": "二、讲授与探究",
                "time_min": max(20, duration // 2),
                "teacher_activity": "分步讲解核心概念，示范关键步骤，组织小组探究。",
                "student_activity": "动手操作/讨论记录，展示交流。",
            },
            {
                "stage": "三、巩固练习",
                "time_min": 10,
                "teacher_activity": "分层练习与即时反馈，针对易错点精讲。",
                "student_activity": "独立完成与互评。",
            },
            {
                "stage": "四、课堂小结与作业",
                "time_min": 8,
                "teacher_activity": "归纳知识结构，布置分层作业与拓展阅读。",
                "student_activity": "自我反思，记录疑问。",
            },
        ],
        "homework": "必做：教材配套练习；选做：拓展阅读或小论文提纲。",
        "board_design": "主板：概念图 + 关键公式/步骤；副板：例题与易错提醒。",
        "reflection_placeholder": "课后填写：目标达成度、学生反馈与改进点。",
    }

    if _has_real_llm():
        prompt = f"""你是资深教研员。请为主题「{topic}」撰写一份完整教案（面向{level}，{grade}），课长{duration}分钟。
教学需求说明：{requirements or "（无则结合课题自拟）"}

请只输出一个 JSON 对象（不要 Markdown），字段如下：
title, meta(课题/课程学段/年级/课型/课时), objectives(三个维度各为字符串数组),
key_points(数组), difficulties(数组),
teaching_process(数组，每项含 stage, time_min, teacher_activity, student_activity),
homework, board_design, reflection_placeholder
内容要具体、可执行，符合中国课堂教案习惯。"""
        raw = llm_client.blocking_chat([{"role": "user", "content": prompt}], max_tokens=4096)
        try:
            js = raw[raw.index("{") : raw.rindex("}") + 1]
            parsed = json.loads(js)
            if isinstance(parsed, dict) and parsed.get("title"):
                return parsed
        except Exception:
            pass
    return skeleton


def write_lesson_docx(lp: Dict[str, Any], out_path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    title = str(lp.get("title") or "教案")
    h0 = doc.add_heading(title, 0)
    h0.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = lp.get("meta") or {}
    if isinstance(meta, dict):
        doc.add_paragraph("一、基本信息")
        for k, v in meta.items():
            doc.add_paragraph(f"{k}：{v}", style="List Bullet")

    obj = lp.get("objectives") or {}
    if isinstance(obj, dict):
        doc.add_paragraph("二、教学目标")
        for dim, lines in obj.items():
            doc.add_paragraph(str(dim), style="Heading 3")
            if isinstance(lines, list):
                for line in lines:
                    doc.add_paragraph(str(line), style="List Bullet")
            else:
                doc.add_paragraph(str(lines))

    for sec, label in (
        ("key_points", "三、教学重点"),
        ("difficulties", "四、教学难点"),
    ):
        doc.add_paragraph(label)
        val = lp.get(sec) or []
        if isinstance(val, list):
            for x in val:
                doc.add_paragraph(str(x), style="List Number")
        else:
            doc.add_paragraph(str(val))

    tp = lp.get("teaching_process") or []
    doc.add_paragraph("五、教学过程")
    if isinstance(tp, list):
        for block in tp:
            if not isinstance(block, dict):
                continue
            doc.add_paragraph(str(block.get("stage", "")), style="Heading 3")
            doc.add_paragraph(f"时长（约）：{block.get('time_min', '')} 分钟")
            doc.add_paragraph("教师活动：" + str(block.get("teacher_activity", "")))
            doc.add_paragraph("学生活动：" + str(block.get("student_activity", "")))

    doc.add_paragraph("六、作业布置")
    doc.add_paragraph(str(lp.get("homework", "")))

    doc.add_paragraph("七、板书设计")
    doc.add_paragraph(str(lp.get("board_design", "")))

    doc.add_paragraph("八、教学反思（课后填写）")
    doc.add_paragraph(str(lp.get("reflection_placeholder", "")))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
