"""模板智能识别服务 - 自动识别上传的PPTX/DOCX文件是PPT模板还是教案模板"""
import os
import re
from typing import Dict, Optional


class TemplateAnalyzer:
    """模板智能识别器"""

    # PPT模板特征关键词
    PPT_KEYWORDS = [
        "封面", "目录", "知识点", "案例", "总结", "互动", "练习",
        "课件", "幻灯片", "讲义", "PPT", "演示", "展示",
        "标题", "副标题", "页眉", "页脚",
    ]

    # 教案模板特征关键词
    LESSON_PLAN_KEYWORDS = [
        "教案", "教学设计", "教学目标", "教学过程", "教学反思",
        "课程标准", "学情分析", "教学重点", "教学难点",
        "导入环节", "新授环节", "巩固练习", "课堂小结",
        "板书设计", "作业布置", "教学评价", "课后反思",
        "教学方法", "教学手段", "教学准备",
        "三维目标", "核心素养", "学科素养",
    ]

    # 教案结构标题模式
    LESSON_PLAN_HEADING_PATTERNS = [
        r"[一二三四五六七八九十]+[、.．]",  # 一、二、三、
        r"第[一二三四五六七八九十\d]+[部节课环]",  # 第一部分、第一节
        r"[\(（][一二三四五六七八九十\d]+[\)）]",  # （一）（二）
    ]

    def analyze_file(self, file_path: str) -> Dict:
        """分析文件，返回模板类型和特征信息"""
        file_ext = os.path.splitext(file_path)[1].lower()

        if file_ext == ".pptx":
            return self._analyze_pptx(file_path)
        elif file_ext in (".docx", ".doc"):
            return self._analyze_docx(file_path)
        else:
            return {
                "template_type": "unknown",
                "confidence": 0,
                "details": "不支持的文件类型",
            }

    def _analyze_pptx(self, file_path: str) -> Dict:
        """分析PPTX文件"""
        try:
            from pptx import Presentation

            prs = Presentation(file_path)
            slide_count = len(prs.slides)

            # 提取所有文本和分析结构
            all_text = ""
            slide_types = []
            has_placeholder_content = False
            empty_slides = 0

            for i, slide in enumerate(prs.slides):
                slide_text = ""
                shape_count = 0
                has_image = False

                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        slide_text += shape.text + "\n"
                    shape_count += 1
                    # 检查是否有图片
                    if shape.shape_type is not None and shape.shape_type == 13:
                        has_image = True

                all_text += slide_text

                # 判断幻灯片类型
                text_lower = slide_text.lower().strip()
                if not text_lower:
                    empty_slides += 1
                    slide_types.append("empty")
                elif i == 0 and len(slide_text.strip()) < 100:
                    slide_types.append("cover")
                elif any(kw in text_lower for kw in ["目录", "目次", "contents", "outline"]):
                    slide_types.append("toc")
                elif any(kw in text_lower for kw in ["总结", "小结", "summary", "回顾"]):
                    slide_types.append("summary")
                elif any(kw in text_lower for kw in ["练习", "习题", "exercise", "互动"]):
                    slide_types.append("exercise")
                elif any(kw in text_lower for kw in ["感谢", "谢谢", "thank"]):
                    slide_types.append("thanks")
                else:
                    slide_types.append("content")

                # 检查是否有占位符文本（模板特征）
                placeholder_markers = [
                    "请输入", "在此输入", "此处输入", "点击输入",
                    "请填写", "在此填写", "添加标题", "添加文本",
                    "click to add", "insert text", "your title",
                ]
                if any(marker in text_lower for marker in placeholder_markers):
                    has_placeholder_content = True

            # 计算PPT模板得分
            ppt_score = 0

            # 有多页幻灯片
            if slide_count >= 3:
                ppt_score += 20
            if slide_count >= 8:
                ppt_score += 10

            # 有封面和目录
            if "cover" in slide_types:
                ppt_score += 15
            if "toc" in slide_types:
                ppt_score += 15

            # 有结构化的页面类型分布
            content_pages = slide_types.count("content")
            if content_pages >= 3:
                ppt_score += 15

            # 有占位符文本（强烈模板信号）
            if has_placeholder_content:
                ppt_score += 25

            # 空页面比例高（模板特征）
            if slide_count > 0 and empty_slides / slide_count > 0.3:
                ppt_score += 15

            # PPT关键词匹配
            ppt_keyword_count = sum(1 for kw in self.PPT_KEYWORDS if kw in all_text)
            ppt_score += min(ppt_keyword_count * 3, 15)

            # 检查是否更像教案(PPTX格式的教案)
            lesson_keyword_count = sum(1 for kw in self.LESSON_PLAN_KEYWORDS if kw in all_text)

            # 判断模板类型
            if lesson_keyword_count >= 5 and lesson_keyword_count > ppt_keyword_count:
                return {
                    "template_type": "lesson_plan_template",
                    "file_format": "pptx",
                    "confidence": min(lesson_keyword_count * 10, 90),
                    "slide_count": slide_count,
                    "slide_types": slide_types,
                    "details": f"检测到教案模板(PPTX格式)，包含{slide_count}页，"
                               f"匹配{lesson_keyword_count}个教案关键词",
                    "suggested_template": self._suggest_lesson_plan_type(all_text),
                }
            else:
                confidence = min(ppt_score, 95)
                return {
                    "template_type": "ppt_template",
                    "file_format": "pptx",
                    "confidence": confidence,
                    "slide_count": slide_count,
                    "slide_types": slide_types,
                    "has_placeholder": has_placeholder_content,
                    "details": f"检测到PPT课件模板，包含{slide_count}页，"
                               f"结构: {self._summarize_slide_types(slide_types)}",
                    "suggested_template": self._suggest_ppt_template(slide_types, all_text),
                }

        except Exception as e:
            return {
                "template_type": "ppt_template",
                "file_format": "pptx",
                "confidence": 30,
                "details": f"文件分析部分失败: {str(e)}，默认识别为PPT模板",
                "suggested_template": "standard",
            }

    def _analyze_docx(self, file_path: str) -> Dict:
        """分析DOCX文件"""
        try:
            from docx import Document

            doc = Document(file_path)

            # 提取所有文本和标题结构
            all_text = ""
            headings = []
            paragraph_count = 0
            table_count = len(doc.tables)

            for para in doc.paragraphs:
                all_text += para.text + "\n"
                paragraph_count += 1

                # 收集标题
                if para.style and para.style.name and "Heading" in para.style.name:
                    headings.append(para.text)
                elif para.text.strip():
                    # 检查是否是格式化的标题（如 一、xxx）
                    for pattern in self.LESSON_PLAN_HEADING_PATTERNS:
                        if re.match(pattern, para.text.strip()):
                            headings.append(para.text.strip())
                            break

            # 计算教案得分
            lesson_score = 0

            # 教案关键词匹配
            lesson_keyword_count = sum(
                1 for kw in self.LESSON_PLAN_KEYWORDS if kw in all_text
            )
            lesson_score += min(lesson_keyword_count * 8, 50)

            # 有结构化标题
            if len(headings) >= 3:
                lesson_score += 20
            if len(headings) >= 6:
                lesson_score += 10

            # 有表格（教案常用）
            if table_count >= 1:
                lesson_score += 10

            # 段落数量适中
            if 20 <= paragraph_count <= 200:
                lesson_score += 10

            # PPT关键词匹配（对比）
            ppt_keyword_count = sum(1 for kw in self.PPT_KEYWORDS if kw in all_text)

            # 判断
            if lesson_score >= 30 and lesson_keyword_count > ppt_keyword_count:
                return {
                    "template_type": "lesson_plan_template",
                    "file_format": "docx",
                    "confidence": min(lesson_score, 95),
                    "heading_count": len(headings),
                    "headings": headings[:10],
                    "table_count": table_count,
                    "details": f"检测到教案模板，包含{len(headings)}个标题、"
                               f"{table_count}个表格、{paragraph_count}个段落",
                    "suggested_template": self._suggest_lesson_plan_type(all_text),
                }
            else:
                return {
                    "template_type": "document_template",
                    "file_format": "docx",
                    "confidence": max(40, 100 - lesson_score),
                    "heading_count": len(headings),
                    "details": f"检测到文档模板，包含{len(headings)}个标题、"
                               f"{paragraph_count}个段落",
                    "suggested_template": "standard",
                }

        except Exception as e:
            return {
                "template_type": "document_template",
                "file_format": "docx",
                "confidence": 30,
                "details": f"文件分析部分失败: {str(e)}",
                "suggested_template": "standard",
            }

    def _summarize_slide_types(self, slide_types: list) -> str:
        """汇总幻灯片类型分布"""
        from collections import Counter
        counts = Counter(slide_types)
        parts = []
        name_map = {
            "cover": "封面",
            "toc": "目录",
            "content": "内容",
            "exercise": "练习",
            "summary": "总结",
            "thanks": "感谢",
            "empty": "空白",
        }
        for t, c in counts.most_common():
            parts.append(f"{name_map.get(t, t)}{c}页")
        return "、".join(parts)

    def _suggest_ppt_template(self, slide_types: list, text: str) -> str:
        """根据PPT结构推荐最匹配的模板"""
        text_lower = text.lower()

        # 检查特征关键词
        if any(kw in text_lower for kw in ["实验", "实验目的", "实验步骤", "实验原理"]):
            return "experimental"
        if any(kw in text_lower for kw in ["研究", "论文", "学术", "文献"]):
            return "lecture"
        if any(kw in text_lower for kw in ["互动", "小组", "讨论", "协作"]):
            return "interactive"

        # 根据页面类型分布
        exercise_count = slide_types.count("exercise")
        if exercise_count >= 2:
            return "interactive"

        return "standard"

    def _suggest_lesson_plan_type(self, text: str) -> str:
        """根据教案内容推荐类型"""
        text_lower = text.lower()

        if any(kw in text_lower for kw in ["实验", "实验报告"]):
            return "experimental"
        if any(kw in text_lower for kw in ["互动", "小组合作", "讨论"]):
            return "interactive"
        if any(kw in text_lower for kw in ["讲座", "学术"]):
            return "lecture"

        return "standard"


# 单例
template_analyzer = TemplateAnalyzer()
