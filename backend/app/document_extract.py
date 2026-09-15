"""Extract plain text from uploaded files (fast path first, then fallbacks)."""
from pathlib import Path
from typing import Iterable, List

# 控制解析量：避免超大 PDF 阻塞请求
MAX_PDF_PAGES = 60
MAX_TOTAL_CHARS = 400_000
MAX_REFERENCE_CHARS = 24_000


def extract_text(path: Path) -> str:
    suf = path.suffix.lower()
    if suf == ".pdf":
        return _pdf(path)
    if suf in (".docx", ".doc"):
        return _docx(path)
    if suf in (".txt", ".md"):
        return _cap(_read_utf8(path))
    if suf in (".pptx", ".ppt"):
        return _pptx(path)
    return _cap(_read_utf8(path))


def resolve_uploaded_references(
    references: Iterable[str] | None,
    *,
    max_chars: int = MAX_REFERENCE_CHARS,
) -> List[str]:
    """把上传文件 ID 解析成有来源标记的文本，普通文本保持兼容。"""
    from . import persist

    registry = persist.load_file_registry().get("files") or {}
    resolved: List[str] = []
    remaining = max(0, int(max_chars))
    for raw in references or []:
        value = str(raw or "").strip()
        if not value or remaining <= 0:
            continue
        meta = registry.get(value)
        if meta:
            path = Path(meta.get("path") or "")
            text = extract_text(path) if path.is_file() else ""
            if not text.strip():
                continue
            name = Path(meta.get("filename") or path.name).name
            value = f"[参考资料：{name}]\n{text.strip()}"
        piece = value[:remaining]
        resolved.append(piece)
        remaining -= len(piece)
    return resolved


def _cap(s: str) -> str:
    if len(s) <= MAX_TOTAL_CHARS:
        return s
    return s[:MAX_TOTAL_CHARS] + "\n\n[内容过长，已截断用于知识库]"


def _read_utf8(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        return ""


def _pdf(path: Path) -> str:
    # 1) PyPDF2 通常比 pdfplumber 快一个数量级
    try:
        import PyPDF2

        reader = PyPDF2.PdfReader(str(path))
        n = min(len(reader.pages), MAX_PDF_PAGES)
        parts = []
        for i in range(n):
            try:
                t = reader.pages[i].extract_text() or ""
            except Exception:
                t = ""
            if t:
                parts.append(t)
        out = "\n".join(parts)
        if len(out.strip()) > 50:
            return _cap(out)
    except Exception:
        pass
    # 2) pdfplumber 版式复杂 PDF
    try:
        import pdfplumber

        parts = []
        with pdfplumber.open(str(path)) as pdf:
            for i, page in enumerate(pdf.pages[:MAX_PDF_PAGES]):
                try:
                    t = page.extract_text() or ""
                except Exception:
                    t = ""
                if t:
                    parts.append(t)
        return _cap("\n".join(parts))
    except Exception:
        return ""


def _docx(path: Path) -> str:
    try:
        import docx

        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text]
        for table in d.tables:
            for row in table.rows:
                parts.append("\t".join(c.text for c in row.cells))
        return _cap("\n".join(parts))
    except Exception:
        return ""


def _pptx(path: Path) -> str:
    try:
        from pptx import Presentation

        prs = Presentation(str(path))
        parts = []
        for slide in prs.slides[:80]:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    parts.append(shape.text)
        return _cap("\n".join(parts))
    except Exception:
        return ""
