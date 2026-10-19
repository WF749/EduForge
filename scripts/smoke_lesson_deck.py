"""课件契约与校验器的回归测试。

钉死两件事：
1. **合格样例必须判合格**（12 份黄金样例不能有误杀）
2. **不合格的必须拦住** —— 尤其是我们**自己此前真实产出过**的那几类：
   - 有占位符（授课教师：XXX / 教材PXX页）
   - 缺教学闭环（只有知识点堆砌，没有练习/小结/作业）
   - 练习只出题不给答案
   - 只有活动名没有活动内容（"小组讨论"但没有讨论题/分工/反馈）

还有一个必须钉的：**不得误杀** —— 英语/语文题里的短下划线是合法填空线，
不是待填占位。误报一多，这个校验就没人信了。

运行：PYTHONPATH=backend uv run --project backend python scripts/smoke_lesson_deck.py
"""

from __future__ import annotations

import copy
import json
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "backend"))

PASS, FAIL = [], []


def check(name, cond, extra=""):
    (PASS if cond else FAIL).append(name)
    print(f"  [{'PASS' if cond else 'FAIL'}] {name}" + (f"  {extra}" if extra else ""))


def main() -> int:
    print("=" * 64)
    print("课件契约与校验 —— 行为测试")
    print("=" * 64)

    from app.services.lesson_deck import (load_golden, validate_deck, find_golden,
                                          find_placeholders, PAGE_SPEC, CLOSURE_STAGES)

    # ---- 1. 黄金样例全部合格 ----
    print("\n== 1. 12 份黄金样例必须判合格（不能有误杀） ==")
    gs = load_golden()
    check("黄金样例已入库", len(gs) == 12, f"{len(gs)} 份")
    bad = []
    for d in gs:
        r = validate_deck(d)
        if not r["ok"]:
            bad.append(f"{d.get('subject')}/{d.get('title')}: {r['errors'][:1]}")
    check("全部合格", not bad, "; ".join(bad[:2]))

    # ---- 2. 反例：占位符必须拦住 ----
    print("\n== 2. 反例：占位符（我们自己真实产出过） ==")
    base = copy.deepcopy(gs[0])
    for ph_text in ["授课教师：XXX", "教材PXX页第3题", "（请根据实际情况补充）",
                    "此处插入图片", "示例内容", "老师自行设计"]:
        d = copy.deepcopy(base)
        d["slides"][1]["items"] = [ph_text, "第二条"]
        r = validate_deck(d)
        check(f"拦住「{ph_text}」", not r["ok"] and any("占位符" in e for e in r["errors"]))

    # 长串下划线 = 留白待填，要拦
    d = copy.deepcopy(base)
    d["slides"][1]["items"] = ["授课教师 ________", "第二条"]
    check("拦住长下划线（留白待填）", not validate_deck(d)["ok"])

    # ---- 3. 不得误杀：题目里的短下划线 ----
    print("\n== 3. 不得误杀：题目里的合法填空线 ==")
    check("find_placeholders 在正常文本放行短下划线",
          find_placeholders("My mother is in the ____.", in_question=True) == [])
    check("但题目里的 TODO 仍要抓",
          find_placeholders("翻译：TODO", in_question=True) != [])
    d = copy.deepcopy(base)
    ex_idx = [i for i, s in enumerate(d["slides"]) if s.get("type") == "exercise"]
    if ex_idx:
        d["slides"][ex_idx[0]]["items"] = [{"q": "My mother is in the ____.", "a": "kitchen"}]
        r = validate_deck(d)
        check("含填空线的练习页仍合格", r["ok"], str(r["errors"][:1]))

    # ---- 4. 反例：缺教学闭环 ----
    print("\n== 4. 反例：缺教学闭环 ==")
    for drop in ("exercise", "summary", "homework"):
        d = copy.deepcopy(base)
        d["slides"] = [s for s in d["slides"] if s.get("type") != drop]
        r = validate_deck(d)
        check(f"去掉 {drop} 页 → 判不合格", not r["ok"], f"缺{r['missing_stages']}")

    # ---- 5. 反例：练习只出题不给答案 ----
    print("\n== 5. 反例：练习必须有答案 ==")
    d = copy.deepcopy(base)
    ex = [s for s in d["slides"] if s.get("type") == "exercise"]
    if ex:
        ex[0]["items"] = [{"q": "物体浸没后浮力是否变化？"}]
        r = validate_deck(d)
        check("练习缺答案 → 拦住", any("缺少答案" in e for e in r["errors"]))

    # ---- 6. 反例：只有活动名没有活动内容 ----
    print("\n== 6. 反例：只给活动名不给内容 ==")
    d = copy.deepcopy(base)
    act = [s for s in d["slides"] if s.get("type") == "activity"]
    if act:
        act[0]["steps"] = []          # 只有"小组讨论"，没有讨论题/分工/反馈
        r = validate_deck(d)
        check("活动 steps 为空 → 拦住", any("为空列表" in e for e in r["errors"]))
    d2 = copy.deepcopy(base)
    d2["slides"] = [{"type": "activity", "title": "小组讨论"}]   # 只有名字
    check("只有活动名 → 拦住", not validate_deck(d2)["ok"])

    # ---- 7. 契约完整性 ----
    print("\n== 7. 契约完整性 ==")
    check("12 种页面类型齐全", len(PAGE_SPEC) == 12, str(sorted(PAGE_SPEC)))
    check("闭环 6 段齐全", len(CLOSURE_STAGES) == 6)
    check("未知类型会被拦", not validate_deck(
        {"meta": {"topic": "x", "subject": "物理"},
         "slides": [{"type": "不存在的类型"}]})["ok"])

    # ---- 8. 查样例 ----
    print("\n== 8. 黄金样例检索 ==")
    check("按学科+课题命中", (find_golden("物理", "浮力") or {}).get("title") == "浮力")
    check("学科不符不命中", find_golden("语文", "浮力") is None)


    # ---- 9. 学科不确定时仍要能命中黄金样例 ----
    print("\n== 9. 学科不确定不能导致检索失败 ==")
    # 实测事故：学科=「综合课程」（界面默认值）时命中不了，学科=空反而能命中。
    # 用默认值惩罚用户是错的。
    for subj in ["综合课程", "", "未设学科", None, "物理"]:
        check(f"学科={subj or '(空)'} 仍命中浮力",
              (find_golden(subj, "浮力") or {}).get("title") == "浮力")
    check("课题不符仍不命中", find_golden("综合课程", "不存在课题") is None)

    # ---- 10. 兜底模板不得把表单字段当知识点 ----
    print("\n== 10. 兜底模板不能产生「XX的定义与本质」垃圾 ==")
    from app.ppt_tasks import _looks_like_meta_field, _build_detailed_content
    for meta in ["内容难度：中等难度", "难度：中等", "时长：45分钟", "课时：2课时",
                 "学科：物理", "课型：新授课", "受众：八年级"]:
        check(f"识别「{meta}」为表单字段", _looks_like_meta_field(meta))
    for good in ["浮力", "阿基米德原理", "细胞的结构", "一元二次方程"]:
        check(f"不误判「{good}」", not _looks_like_meta_field(good))
    # 用带字段的 outline 走一遍，确认生成的页不再含 "内容难度：中等难度的定义"
    slides = _build_detailed_content("浮力", "内容难度：中等难度\n浮力产生的原因", "八年级")
    joined = json.dumps(slides, ensure_ascii=False)
    check("兜底内容不再出现「内容难度：中等难度的定义」",
          "内容难度：中等难度的定义" not in joined)
    check("兜底内容退回用课题本身", "浮力" in joined)


    # ---- 11. 渲染后不得有文字溢出（本次事故的直接防线） ----
    print("\n== 11. 渲染溢出检测 ==")
    from app.services.deck_renderer import render_deck
    from pptx import Presentation
    import tempfile

    def _scan_overflow(path):
        """估算每个文本框需要的高度 vs 实际高度。"""
        pr = Presentation(path)
        bad = []
        for i, sl in enumerate(pr.slides):
            for sh in sl.shapes:
                if not sh.has_text_frame:
                    continue
                txt = sh.text_frame.text
                if len(txt) < 20:
                    continue
                size = next((r.font.size.pt for pa in sh.text_frame.paragraphs
                             for r in pa.runs if r.font.size), None)
                if not size:
                    continue
                w_px, h_px = sh.width / 9525, sh.height / 9525
                per_line = max(1, int(w_px / (size * 1.34)))
                lines = sum(max(1, (len(pa.text) + per_line - 1) // per_line)
                            for pa in sh.text_frame.paragraphs)
                if lines * size * 2.15 > h_px * 1.05:
                    bad.append((i + 1, round(lines * size * 2.15), round(h_px)))
        return bad

    for band, gidx, label in [("初中", 6, "浮力"), ("小学", 0, "富饶的西沙群岛"),
                              ("高中", 11, "氧化还原反应")]:
        d = load_golden()[gidx]
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "t.pptx")
            render_deck(d, out)
            bad = _scan_overflow(out)
            check(f"{label}（{band}）无文字溢出", len(bad) == 0,
                  f"{len(bad)} 处 {bad[:2]}")

    # 极端内容：超长题干 + 超长答案，靠截断兜住
    stress = json.loads(json.dumps(load_golden()[6], ensure_ascii=False))
    for sl in stress["slides"]:
        if sl.get("type") == "exercise":
            sl["items"] = [{"q": "很长的题干" * 30, "a": "很长的解析" * 40} for _ in range(3)]
        if sl.get("type") == "example":
            sl["answer"] = "很长的答案" * 40
            sl["problem"] = "很长的题目" * 20
    with tempfile.TemporaryDirectory() as tmp:
        out = os.path.join(tmp, "stress.pptx")
        render_deck(stress, out)
        bad = _scan_overflow(out)
        check("超长内容也不溢出（靠截断兜住）", len(bad) == 0, f"{len(bad)} 处 {bad[:2]}")

    print("\n" + "=" * 64)
    print(f"通过 {len(PASS)} 项，失败 {len(FAIL)} 项")
    if FAIL:
        for n in FAIL:
            print(f"  - {n}")
    print("=" * 64)
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
