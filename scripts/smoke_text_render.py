"""成品文本渲染闸门的回归测试。

**为什么单独一个脚本**：`smoke_lesson_docx.py` 用的是干净的假数据，35 项全绿，
但真实模型返回的形状（被截断的 repr、值里含引号的字面量、多行里混着一行 repr）
一条都没覆盖 —— 线上照样产出 `{'具体目标': "…"` 和 `['控制变量法验证：…']`。

本脚本的用例**全部来自真实事故现场**：
  - 教案 Word 表格里的 `1. {'具体目标': "能准确说出浮力是…"}`
  - PPT 第 6 页要点里的 `['控制变量法验证：①…；②…']`
  - PPT 第 7 页的 `['密度判据：ρ_物＜ρ_液→上浮终漂浮']`

运行：PYTHONPATH=backend uv run --project backend python scripts/smoke_text_render.py
不消耗真实模型额度。
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

PASS = []
FAIL = []


def check(name: str, cond: bool, extra: str = "") -> None:
    (PASS if cond else FAIL).append(name)
    print(f"  [{'PASS' if cond else 'FAIL'}] {name}" + (f"  {extra}" if extra else ""))


def main() -> int:
    print("=" * 64)
    print("文本渲染闸门 —— 事故形态回归")
    print("=" * 64)

    from app.services import text_render as tr

    # ---------------- 1. 教案事故形态 ----------------
    print("\n== 1. 教案 Word 里的真实事故形态 ==")
    cases = [
        ("教学目标 dict",
         "{'具体目标': \"能准确说出浮力是液体对物体上下表面压力差的合力\"}"),
        ("检测题 dict",
         "{'题目': '体积为200 cm³的实心铁块浸没在水中，求所受浮力。', '评分标准': '①单位换算正确（200 cm³=2×10⁻⁴ m³）'}"),
        ("维度+内容 dict",
         "{'维度': '物理观念', '内容': ['能说出浮力定义', '能画出受力图']}"),
        ("嵌套 dict",
         "{'a': {'b': '深层值'}, 'c': 1}"),
        ("带行首标记",
         "1. {'具体目标': '能复述阿基米德原理'}"),
    ]
    for name, raw in cases:
        out = tr.safe_text(raw)
        check(f"{name} → 无 repr 残留", not tr.looks_like_repr(out), repr(out[:52]))
        check(f"{name} → 内容未丢失", len(out.strip()) > 4, f"len={len(out)}")

    # ---------------- 2. PPT 事故形态 ----------------
    print("\n== 2. PPT 要点里的真实事故形态 ==")
    ppt_cases = [
        ("第6页 控制变量法",
         "• 公式含义：浮力大小等于被物体排开的那部分液体所受的重力\n"
         "• 关键变量只有两个：ρ_液（液体密度）和V_排\n"
         "• ['控制变量法验证：①同种液体中，改变浸入体积→F_浮线性变化；②相同V_排下换液体→F_浮改变']"),
        ("第7页 密度判据",
         "• 动态过程：F_浮＞G_物→上浮\n"
         "• ['密度判据：ρ_物＜ρ_液→上浮终漂浮；ρ_物＝ρ_液→悬浮']"),
        ("第9页 对比强调",
         "• 轮船：空心结构→增大V_排\n"
         "• ['对比强调：轮船靠调V_排（被动适应），潜水艇靠调G_物（主动控制）']"),
    ]
    for name, raw in ppt_cases:
        out = tr.safe_text(raw)
        check(f"{name} → 无 repr 残留", not tr.looks_like_repr(out), repr(out[:52]))
        # 多行输入应当保留多行结构，而不是被压成一句
        check(f"{name} → 保留要点分行", out.count("\n") >= 1, f"{out.count(chr(10))} 行")

    # ---------------- 3. 解析失败的降级 ----------------
    print("\n== 3. literal_eval 失败时必须降级，不能原样保留 ==")
    broken = [
        "{'具体目标': '能说出浮力产生的原因（截断",
        "{'题目': \"含'单引号'的题目\", '评分标准': 'ok'}",
        "['要点一', '要点二",
    ]
    for raw in broken:
        out = tr.safe_text(raw)
        check(f"截断/引号异常仍有输出：{raw[:24]}…", len(out.strip()) > 0)
        # 降级后不应该还是原样的一整坨
        check(f"  未原样保留：{raw[:24]}…",
              out.strip() != raw.strip(), repr(out[:46]))

    # ---------------- 4. 正常内容不能被破坏 ----------------
    print("\n== 4. 正常内容不能被误伤 ==")
    normals = [
        "能准确说出浮力的定义",
        "• 液体内部存在压强\n• 随深度增加而增大",
        "F_浮 = ρ_液 g V_排",
        "",
    ]
    for raw in normals:
        out = tr.safe_text(raw)
        check(f"原样保留：{raw[:22] or '(空)'!r}", out.replace("• ", "").strip() in raw or raw == "" or out == raw,
              repr(out[:46]))

    # ---------------- 5. 编号块 ----------------
    print("\n== 5. 编号块 ==")
    out = tr.safe_text(["能说出浮力定义", "能画出受力图"], numbered=True)
    check("列表 → 带序号", out.startswith("1. ") and "\n2. " in out, repr(out[:40]))
    check("列表 → 无方括号", "[" not in out and "]" not in out, repr(out[:40]))

    # ---------------- 6. PPT 端到端 ----------------
    print("\n== 6. PPT 渲染端到端（含泄漏形态） ==")
    try:
        from pptx import Presentation
        from app.ppt_tasks import _write_pptx
    except Exception as e:  # 缺依赖时降级：只测文本层
        print(f"  （跳过：{e}）")
        Presentation = None
    if Presentation is not None:
        import tempfile
        slides = [
            {"page_type": "cover", "main_title": "浮力",
             "subtitle": "面向 九年级 ｜ 物理课程"},
            {"page_type": "toc", "main_title": "课程目录",
             "bullet_points": "学习目标 · 核心概念 · 原理讲解"},
            {"page_type": "content", "main_title": "浮力产生的根本原因",
             "bullet_points": "• 液体内部存在压强\n• ['控制变量法验证：①改变V_排']"},
        ]
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "t.pptx"
            _write_pptx(p, "浮力", slides, "蓝色系")
            prs = Presentation(p)
            check("16:9 宽屏", prs.slide_width > 12_000_000, f"width={prs.slide_width}")
            check("每页有设计的形状（>2）",
                  all(len(s.shapes) > 2 for s in prs.slides),
                  str([len(s.shapes) for s in prs.slides]))
            texts = [sh.text_frame.text for s in prs.slides for sh in s.shapes
                     if sh.has_text_frame]
            check("成品无 repr", not any(tr.looks_like_repr(t) for t in texts))

    # ---------------- 7. 配图链路（注入假生图，绝不真调接口烧额度） ----------------
    print("\n== 7. 配图链路：三级降级 ==")
    try:
        import os
        from app import ppt_tasks

        slides = [
            {"page_type": "cover", "main_title": "浮力", "bullet_points": ""},
            {"page_type": "content", "main_title": "浮力产生的原因",
             "bullet_points": "• 液体内部存在压强\n• 压力差竖直向上"},
            {"page_type": "toc", "main_title": "课程目录",
             "bullet_points": "目标 · 概念 · 练习"},
        ]

        # 7.1 关闭配图
        os.environ["EDUFORGE_PPT_IMAGE_MODE"] = "off"
        check("模式 off → 不配图", ppt_tasks.prepare_images(
            "浮力", slides, "蓝色系", "物理", "fake-key") == {})

        # 7.2 AI 成功
        os.environ["EDUFORGE_PPT_IMAGE_MODE"] = "ai"
        from app.services import image_gen
        orig = image_gen.generate
        image_gen.generate = lambda p, k, **kw: "C:/fake/ai.png"
        try:
            got = ppt_tasks.prepare_images("浮力", slides, "蓝色系", "物理", "fake-key")
            check("AI 生图成功 → 返回图片路径", len(got) >= 1 and "ai.png" in list(got.values())[0])
            # 目录页不该被配图（插不进去，白白烧钱）
            lays = ppt_tasks.layouts_for(slides)
            check("目录页不配图", 2 not in got, f"版式={lays}")
        finally:
            image_gen.generate = orig

        # 7.3 AI 失败 → 降级本地装饰图（且不能 segfault）
        image_gen.generate = lambda p, k, **kw: None
        try:
            got = ppt_tasks.prepare_images("浮力", slides, "蓝色系", "物理", "fake-key")
            check("AI 失败 → 降级为本地装饰图", len(got) >= 1, str(list(got.values())[:1]))
            check("降级产物是真实文件", all(os.path.exists(v) for v in got.values()))
        finally:
            image_gen.generate = orig

        # 7.4 生图异常不能让主流程崩
        def _boom(p, k, **kw):
            raise RuntimeError("接口炸了")
        image_gen.generate = _boom
        try:
            got = ppt_tasks.prepare_images("浮力", slides, "蓝色系", "物理", "fake-key")
            check("生图抛异常 → 仍不中断", isinstance(got, dict))
        finally:
            image_gen.generate = orig

        # 7.5 提示词要能约束"图里不要出现文字"
        pr = image_gen.build_prompt("浮力", "物理", "阿基米德原理")
        check("提示词含主题与学科", "浮力" in pr and "物理" in pr)
        check("提示词明确禁止图上文字", "文字" in pr and "水印" in pr)
        os.environ.pop("EDUFORGE_PPT_IMAGE_MODE", None)
    except Exception as e:
        check(f"配图链路测试异常：{e}", False)

    print("\n" + "=" * 64)
    print(f"通过 {len(PASS)} 项，失败 {len(FAIL)} 项")
    if FAIL:
        for n in FAIL:
            print(f"  - {n}")
    print("=" * 64)
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
