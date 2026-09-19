"""教案 Word 排版的冒烟测试（真建一份 docx 再回头验，不联网、秒级）。

存在的理由：文档排版的问题**在界面上看不出来**——教师下载完、打印出来、
或者在 Word 里翻到第二页才发现。而且它坏的形态都很安静：

  1. python-docx 的默认模板是 **Letter**（21.6×27.9cm）不是 A4。不显式设置的话
     打印机按 Letter 排，右边距对不上、最后一列被裁 —— 而"生成成功"照样返回。
  2. 渲染代码里到处是 `str(v)`，模型一旦把一条目标写成 `{"具体目标": "…"}`，
     **Python 的 repr 就会原样印进 Word**。线上真的出现过：
     教学目标表格里躺着 `{'具体目标': "能准确说出…`。
  3. `- ` 开头的伪列表在 Word 里就是个破折号，没有悬挂缩进；
     多行缩进只能靠空格硬凑，文档里真的出现过 `-    单位：ρ→kg/m³`。

所以这个测试不只是"跑通"，而是拿三种**真实出现过的坏输入**去建文档，
再反过来检查产物里没有这些痕迹。
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from docx import Document  # noqa: E402
from docx.table import Table  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from app.services.llm_lesson_plan_all_stages import LessonPlanGenerator  # noqa: E402

PASS: list = []
FAIL: list = []


def check(name: str, cond: bool, extra: str = "") -> None:
    (PASS if cond else FAIL).append(name)
    print(f"  [{'PASS' if cond else 'FAIL'}] {name}" + (f"  {extra}" if extra else ""))


# ---------------------------------------------------------------------------
# 三种真实出现过的坏输入：dict 藏在 list 里、list 藏在 dict 里、嵌套两层
# ---------------------------------------------------------------------------
NASTY_PLAN = {
    "教案基本信息": {
        "课题": "浮力", "学科": "物理", "年级": "八年级", "授课时长": "45分钟",
        "教材版本": {"版本": "人教版", "章节": "第十章第1节"},   # 值本身是 dict
    },
    # 就是这一条在线上印成了 `{'具体目标': "能准确说出…`
    "教学目标": [
        {"具体目标": "能准确说出浮力的定义"},
        {"具体目标": "能用 F_浮 = G_排 计算浸没物体的浮力"},
    ],
    "教学重难点": {
        "教学重点": ["浮力产生的压力差本质", "阿基米德原理"],
        "教学难点": {"微观机理": "从液体压强分布推导压力差合力", "认知障碍": "与物重无关"},
    },
    "教学准备": {
        "教师准备": [{"名称": "弹簧测力计", "数量": 2}, "溢水杯", {"名称": "物块", "数量": 3}],
        "学生准备": ["预习课本", "完成前测问卷"],
        "自备材料": ["橡皮泥", "硬币"],   # 未匹配字段应归入教师准备
    },
    "教学过程设计": [
        {"环节": "情境启思", "时长": "5分钟", "教师活动": "播放轮船离港视频",
         "学生活动": "观察并描述现象", "设计意图": "激活认知锚点"},
        {"环节": "实验探因", "时长": "15分钟", "教师活动": "演示弹簧测力计实验",
         "学生活动": "分组读数记录", "设计意图": "量化浮力",
         "补充说明": ["注意读数前调零", "物块要完全浸没"]},   # 额外字段藏 list
    ],
    # 板书里带层级的"伪列表"，原实现会把空格缩进原样输出
    "板书设计/课件框架": {
        "结构化板书要点": ["标题：浮力——谁托起了轮船？", "一、浮力本质：F_浮 = F_向上 − F_向下",
                          "    单位：ρ→kg/m³，V_排→m³"],
        "课件页码一一对应": {"PPT1": "情境启思", "PPT2": "实验探因"},
    },
    "作业布置与课后拓展": {
        "课后作业": {"基础题": ["教材 P90 第 1、3 题", "教材 P90 第 5 题"], "提升题": "比较酒精与水中浮力"},
        "分层学习要求": ["学困生：完成基础题", "中等生：全部完成"],
    },
    "教学反思与应急预案": "预计 90% 学生能正确计算浸没物体浮力。",
}

gen = LessonPlanGenerator()
tmp = tempfile.mkdtemp(prefix="eduforge_docx_")
path = gen._export_to_docx(NASTY_PLAN, tmp, topic="浮力")
doc = Document(path)

# ---- 收集全部文本（正文段落 + 表格单元格，顺序按文档流） ----
para_texts: list[str] = []
cell_texts: list[str] = []
tables: list = []
runs_by_size: dict[float, int] = {}

for child in doc.element.body.iterchildren():
    tag = child.tag.split("}")[-1]
    if tag == "p":
        p = Paragraph(child, doc)
        if p.text.strip():
            para_texts.append(p.text)
        for r in p.runs:
            if r.font.size:
                runs_by_size[r.font.size.pt] = runs_by_size.get(r.font.size.pt, 0) + 1
    elif tag == "tbl":
        tb = Table(child, doc)
        tables.append(tb)
        for row in tb.rows:
            for c in row.cells:
                cell_texts.append(c.text)
                for p in c.paragraphs:
                    for r in p.runs:
                        if r.font.size:
                            runs_by_size[r.font.size.pt] = runs_by_size.get(r.font.size.pt, 0) + 1

all_text = "\n".join(para_texts + cell_texts)

print("\n=== 1. 纸张必须是 A4（默认是 Letter，不设置就会印错）===")
sec = doc.sections[0]
check("宽 21.0cm", abs(sec.page_width.cm - 21.0) < 0.05, f"实得 {sec.page_width.cm:.2f}")
check("高 29.7cm", abs(sec.page_height.cm - 29.7) < 0.05, f"实得 {sec.page_height.cm:.2f}")

print("\n=== 2. 绝不能把 Python 的字面量印进文档（线上事故）===")
check("没有 dict 开括号痕迹", "{'" not in all_text, "出现 {'")
check("没有 list-of-dict 痕迹", "[{" not in all_text, "出现 [{")
check("没有 dict 的键值引号", "': '" not in all_text and '":' not in all_text)
check("没有 None 字面量", "None" not in all_text)
check("没有 False/True 字面量", "True" not in all_text and "False" not in all_text)
check("教学目标里读得出原句", "能准确说出浮力的定义" in all_text)
check("嵌套 dict 被摊成「键：值」", "微观机理：从液体压强分布推导压力差合力" in all_text)
check("list 里的 dict 被摊平", "弹簧测力计" in all_text and "数量" in all_text)
check("未匹配字段没丢", "自备材料" in all_text and "橡皮泥" in all_text)
check("过程表里的额外字段没丢", "补充说明" in all_text and "物块要完全浸没" in all_text)

print("\n=== 3. 不能再有伪列表与空格硬缩进 ===")
cell_lines = [ln for t in cell_texts for ln in t.split("\n")]
bad_dash = [t for t in para_texts + cell_lines if t.strip().startswith("- ")]
check("没有 `- ` 开头的伪列表项", not bad_dash, f"实得 {len(bad_dash)} 条：{bad_dash[:2]}")
bad_space = [t for t in para_texts if t.startswith("    ")]
check("正文段落没有靠空格做缩进", not bad_space, f"实得 {len(bad_space)} 条：{bad_space[:2]}")

print("\n=== 4. 字号层级一致（不再 11/12pt 混着用）===")
sizes = sorted(runs_by_size)
check("正文/表格字号只出现约定的几种",
      set(sizes) <= {22.0, 16.0, 14.0, 12.0, 10.5}, f"实得 {sizes}")
check("表格内是 10.5pt", 10.5 in runs_by_size)
check("章节标题 16pt", 16.0 in runs_by_size, f"实得 {sizes}")
check("主标题 22pt", 22.0 in runs_by_size)
check("不再出现 11pt（旧实现里与 12pt 混用）", 11.0 not in runs_by_size)

print("\n=== 5. 表格布局（跨页表头 / 固定列宽）===")
check("一共建了 5 张表", len(tables) == 5, f"实得 {len(tables)}")

def has_header_repeat(tb) -> bool:
    tr_pr = tb.rows[0]._tr.find(qn("w:trPr"))
    return tr_pr is not None and tr_pr.find(qn("w:tblHeader")) is not None

def layout_of(tb):
    el = tb._tbl.tblPr.find(qn("w:tblLayout"))
    return el.get(qn("w:type")) if el is not None else None

check("教学过程表：表头跨页重复", has_header_repeat(tables[4]))
check("教学目标表：表头跨页重复", has_header_repeat(tables[1]))
check("所有表都关了自动伸缩（否则 5 列会被挤成一团）",
      all(layout_of(tb) == "fixed" for tb in tables),
      f"实得 {[layout_of(tb) for tb in tables]}")

print("\n=== 6. 列表是真的列表（悬挂缩进），不是破折号加空格 ===")
bullets = []
for child in doc.element.body.iterchildren():
    if child.tag.split("}")[-1] == "p":
        p = Paragraph(child, doc)
        pf = p.paragraph_format
        if p.text.strip().startswith(("•", "◦", "▪")):
            bullets.append((p.text.strip()[:24], pf.left_indent.pt if pf.left_indent else 0,
                            pf.first_line_indent.pt if pf.first_line_indent else 0))
check("板书/作业部分渲染出了项目符号", len(bullets) >= 5, f"实得 {len(bullets)} 条")
check("项目符号段落都有悬挂缩进（首行负缩进）",
      bool(bullets) and all(fl < 0 for _, _, fl in bullets),
      f"实得 {[fl for _, _, fl in bullets][:4]}")
check("层级用正缩进表达", bool(bullets) and all(li > 0 for _, li, _ in bullets))
check("换行后不会顶格（左缩进 ≥ 悬挂量）",
      bool(bullets) and all(li + fl >= 0 for _, li, fl in bullets))

print("\n=== 7. 内容层已被污染时的自愈（模型把结构写成字符串）===")
# 线上真实产出：作业与教学目标的值是两个**字符串**，里面是 Python 的 repr
HEAL_PLAN = {
    "教案基本信息": {"课题": "浮力"},
    "教学目标": {"物理观念": ["{'具体目标': '能说出浮力方向', '课标依据': '课标2.2.3'}"]},
    "作业布置与课后拓展": {
        "课后作业": [
            "- {'题目': '体积为200cm³的铁块浸没水中，求浮力', '评分标准': '单位换算1分，公式1分'}",
            "1. {'题目': '木块漂浮，求V_排', '评分指标': 'V_排=60cm³'}",
        ],
    },
    # 这一段是"像字面量但其实不是"的反例：必须原样保留，不能被吃掉
    "教学反思与应急预案": {"备注": "可以用 { } 表示集合，注意与浮力无关"},
}
heal_path = gen._export_to_docx(HEAL_PLAN, tmp, topic="自愈测试")
heal_doc = Document(heal_path)
heal_texts = []
for child in heal_doc.element.body.iterchildren():
    tag = child.tag.split("}")[-1]
    if tag == "p":
        heal_texts.append(Paragraph(child, heal_doc).text)
    elif tag == "tbl":
        for row in Table(child, heal_doc).rows:
            for c in row.cells:
                heal_texts.append(c.text)
healed = "\n".join(heal_texts)
check("repr 字符串被还原成结构", "'具体目标'" not in healed and "'评分标准'" not in healed)
check("还原后读得出原内容", "能说出浮力方向" in healed and "单位换算1分，公式1分" in healed)
check("键变成「键：值」而不是引号", "题目：" in healed and "评分标准：" in healed)
check("行首的 `- ` / `1. ` 也被剥掉", "- {" not in healed and "1. {" not in healed)
check("不像字面量的花括号文本原样保留", "可以用 { } 表示集合" in healed)

print("\n=== 8. 空模块不产生垃圾段落 ===")
empty_plan = {"教案基本信息": {"课题": "空课"}}
p2 = gen._export_to_docx(empty_plan, tmp, topic="空课")
d2 = Document(p2)
texts2 = [x.text for x in d2.paragraphs if x.text.strip()]
check("空教案不崩", bool(texts2))
check("空教案里没有 None 字样", "None" not in "\n".join(texts2))
check("空教案仍有 A4 纸张", abs(d2.sections[0].page_width.cm - 21.0) < 0.05)

print(f"\n通过 {len(PASS)} 项，失败 {len(FAIL)} 项")
if FAIL:
    print("失败项：" + "、".join(FAIL))
    sys.exit(1)
